from docx import Document
from docx2pdf import convert

def generate_word_document(data: dict, template_path: str, output_path: str):
    """Generates a Word document by replacing placeholders with data and converts it to PDF.

    Args:
        data (dict): A dictionary where the key is the placeholder and the value is the replacement text.
        template_path (str): Path to the template Word document.
        output_path (str): Path where the output Word document will be saved.
    """
    # Load the template document
    doc = Document(template_path)
    
    # Replace placeholders with corresponding data
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, str(value))

    # Save the modified document
    doc.save(output_path)

    # Convert the document to PDF
    convert(output_path, output_path.replace(".docx", ".pdf"))

# Example usage
if __name__ == "__main__":
    data = {
        "{placeholder1}": "hello",
        "{placeholder2}": "world",
        # Add more placeholders and their corresponding values
    }
    
    template_path = "templates.docx"  # Specify the path to your Word template
    output_path = "output.docx"      # Specify the path where the output will be saved
    
    try:
        generate_word_document(data, template_path, output_path)
        print(f"Document generated and saved as {output_path.replace('.docx', '.pdf')}")
    except Exception as e:
        print(f"Error: {e}")
