from docx import Document
from docx2pdf import convert

def generate_word_document(data, template_path, output_path):
    # Load the template
    doc = Document(template_path)
    
    # Replace placeholders with data
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, str(value))

    # Save the modified document
    doc.save(output_path)

    # Convert to PDF
    convert(output_path, output_path.replace(".docx", ".pdf"))

# Example usage
if __name__ == "__main__":
    data = {
        "{placeholder1}": "value1",
        "{placeholder2}": "value2",
        # Add more placeholders and their corresponding values
    }
    
    template_path = "template.docx"
    output_path = "output.docx"
    
    try:
        generate_word_document(data, template_path, output_path)
        print(f"Document generated and saved as {output_path.replace('.docx', '.pdf')}")
    except Exception as e:
        print(f"Error: {e}")
